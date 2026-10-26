import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import os
import io
import copy
import json

# Folium (Leaflet) untuk peta choropleth -- opsional, aplikasi tetap jalan tanpanya
try:
    import folium
    from streamlit_folium import st_folium
    FOLIUM_TERSEDIA = True
except ImportError:
    FOLIUM_TERSEDIA = False

# ==========================================
# 1. KONFIGURASI HALAMAN
# ==========================================
st.set_page_config(
    page_title="Peta Segmentasi UMKM DKI Jakarta",
    page_icon="🏪",
    layout="wide",
    initial_sidebar_state="expanded"
)

FEATURES = ['sektor_makro', 'kelompok_tahun_berdiri']

# Palet warna klaster (6 warna, cukup untuk k maksimum kandidat Elbow k=2..6) -- warna mengikuti
# id klaster dan dipakai konsisten di kartu profil, pie chart, dan heatmap. k per wilayah dibaca dari
# 02c_k_optimal.csv (lihat notebooks/02_kmodes_partitioning.ipynb bagian 4), bukan konstanta.
CLUSTER_COLORS = ["#EF4444", "#3B82F6", "#10B981", "#F59E0B", "#8B5CF6", "#EC4899"]
CLUSTER_BG = ["#FEF2F2", "#EFF6FF", "#ECFDF5", "#FFFBEB", "#F5F3FF", "#FDF2F8"]

# Kode wilayah BPS (properti 'CC_2' di GeoJSON) -> nilai kolom 'wilayah' di dataset
KODE_WILAYAH = {
    '3171': 'KOTA ADM. JAKARTA SELATAN',
    '3172': 'KOTA ADM. JAKARTA TIMUR',
    '3173': 'KOTA ADM. JAKARTA PUSAT',
    '3174': 'KOTA ADM. JAKARTA BARAT',
    '3175': 'KOTA ADM. JAKARTA UTARA',
    '3101': 'KAB. ADM. KEP. SERIBU',
}
GEOJSON_FILENAME = "dki_wilayah.geojson"
WILAYAH_KEP_SERIBU = 'KAB. ADM. KEP. SERIBU'

# Ambang ukuran sampel minimum supaya klaster/sektor kecil tidak ikut dilabeli (dipakai label_ocean())
MIN_N = 30
MIN_PCT = 1.0

# Ambang Red/Neutral/Blue Ocean, murni dari Location Quotient (LQ) --
# 0.8/1.25 mengikuti konvensi standar literatur ekonomi regional (Location Quotient), bukan angka buatan sendiri.
LQ_BATAS_BLUE = 0.8
LQ_BATAS_RED = 1.25

BADGE_WARNA = {'blue': '#3B82F6', 'yellow': '#F59E0B', 'red': '#EF4444', 'gray': '#9CA3AF'}
BADGE_EMOJI = {'blue': '🔵', 'yellow': '🟡', 'red': '🔴', 'gray': '⚪'}

# Penjelasan istilah teknis dalam bahasa awam, dipakai lewat tooltip/help= di berbagai widget
LQ_TOOLTIP = "LQ > 1 berarti sektor ini lebih padat di wilayah ini dibanding rata-rata DKI Jakarta; LQ < 1 berarti lebih jarang."
OCEAN_TOOLTIP = "LQ membandingkan porsi sektor ini di satu wilayah terhadap porsi sektor ini di seluruh DKI Jakarta. LQ < 0.8 = kurang terwakili di wilayah ini (Blue Ocean). LQ > 1.25 = lebih padat dari rata-rata (Red Ocean)."

# Kustomisasi CSS Tampilan
st.markdown("""
    <style>
    .main-title {
        font-size: 26px;
        font-weight: 700;
        color: #1E3A8A;
        margin-bottom: 2px;
    }
    .sub-title {
        font-size: 15px;
        color: #4B5563;
        margin-bottom: 20px;
    }
    </style>
""", unsafe_allow_html=True)

# ==========================================
# 2. PEMUATAN DATA (CACHED)
# ==========================================
def _find_path(filename):
    possible_paths = [
        os.path.join("data", "processed", filename),
        os.path.join("..", "data", "processed", filename),
        filename
    ]
    for path in possible_paths:
        if os.path.exists(path):
            return path
    return None

@st.cache_data(show_spinner=False)
def load_kmodes_data():
    """Memuat dataset terklasterisasi (k optimal per wilayah, 2 fitur) dari direktori data."""
    path = _find_path("02_umkm_kmodes_clustered.csv")
    if path is None:
        return None
    df = pd.read_csv(path)
    df['cluster'] = df['cluster'].astype(int)
    return df

@st.cache_data(show_spinner=False)
def load_elbow_data():
    """Memuat ringkasan biaya Elbow (Cost vs k, k=2..6) per wilayah, dari notebooks/02_kmodes_partitioning.ipynb."""
    path = _find_path("02b_elbow_costs.csv")
    if path is None:
        return None
    return pd.read_csv(path)

@st.cache_data(show_spinner=False)
def load_k_optimal():
    """Memuat tabel k optimal per wilayah (aturan titik terjauh dari garis lurus pada kurva
    Elbow: k_optimal, kekuatan_siku, k_pembanding, selisih_d, catatan), dari
    notebooks/02_kmodes_partitioning.ipynb."""
    path = _find_path("02c_k_optimal.csv")
    if path is None:
        return None
    return pd.read_csv(path)

@st.cache_data(show_spinner=False)
def load_geojson():
    """Memuat batas 6 wilayah administrasi DKI Jakarta (data/geo/dki_wilayah.geojson) dan
    mencocokkan tiap fitur ke kolom 'wilayah' lewat properti 'CC_2' (kode wilayah BPS).
    Mengembalikan (geojson, pesan_error); pesan_error None jika semua wilayah cocok."""
    possible_paths = [
        os.path.join("data", "geo", GEOJSON_FILENAME),
        os.path.join("..", "data", "geo", GEOJSON_FILENAME),
    ]
    path = next((p for p in possible_paths if os.path.exists(p)), None)
    if path is None:
        return None, f"Berkas `data/geo/{GEOJSON_FILENAME}` tidak ditemukan."

    try:
        with open(path, encoding="utf-8") as f:
            geojson = json.load(f)
    except (OSError, ValueError) as e:
        return None, f"Berkas `{path}` gagal dibaca sebagai GeoJSON ({e})."

    features = geojson.get("features", [])
    tanpa_kode = [i for i, ft in enumerate(features) if "CC_2" not in ft.get("properties", {})]
    if tanpa_kode:
        return None, f"Properti `CC_2` (kode wilayah BPS) tidak ditemukan pada {len(tanpa_kode)} fitur GeoJSON."

    kode_asing = sorted({str(ft["properties"]["CC_2"]) for ft in features} - set(KODE_WILAYAH))
    if kode_asing:
        return None, f"Kode `CC_2` tidak dikenali: {', '.join(kode_asing)} (yang diharapkan: {', '.join(KODE_WILAYAH)})."

    for ft in features:
        ft["properties"]["wilayah"] = KODE_WILAYAH[str(ft["properties"]["CC_2"])]

    hilang = sorted(set(KODE_WILAYAH.values()) - {ft["properties"]["wilayah"] for ft in features})
    if hilang:
        return None, f"Wilayah berikut tidak ada di GeoJSON: {', '.join(hilang)}."

    return geojson, None

@st.cache_data(show_spinner=False)
def load_evaluation_summary():
    """Memuat tabel profil klaster final (k optimal per wilayah) -- kategori dominan
    deskriptif, Location Quotient (LQ), dan label Red/Neutral/Blue Ocean per
    klaster -- dari notebooks/03_interpretasi_dan_penerapan_ocean.ipynb."""
    path = _find_path("03_kmodes_evaluation_summary.csv")
    if path is None:
        return None
    return pd.read_csv(path)

# ==========================================
# 3. RED/NEUTRAL/BLUE OCEAN (murni Location Quotient)
# ==========================================
# Murni berbasis frekuensi/rasio -- tidak memakai ukuran jarak maupun komponen
# pertumbuhan usaha apa pun. Ambang 0.8/1.25 mengikuti konvensi Location Quotient
# di literatur ekonomi regional, bukan hasil pembobotan buatan sendiri.

def hitung_referensi_lq(d, sektor_col='sektor_makro'):
    """Distribusi sektor se-DKI Jakarta, dihitung sekali dari seluruh data (semua
    wilayah) dan dipakai sebagai pembanding/baseline untuk semua perhitungan LQ."""
    return d[sektor_col].value_counts(normalize=True)

def hitung_lq(sub_wilayah, sektor, referensi_dki, sektor_col='sektor_makro'):
    """
    Location Quotient: (share sektor di wilayah ini) / (share sektor sama se-DKI Jakarta).
    sub_wilayah   : subset data 1 wilayah, SEMUA sektor (bukan yang sudah difilter ke 1 sektor).
    referensi_dki : hasil hitung_referensi_lq(), dihitung sekali di luar loop wilayah/klaster.
    """
    share_lokal = (sub_wilayah[sektor_col] == sektor).mean()
    share_referensi = referensi_dki.get(sektor, None)
    if share_referensi is None or share_referensi == 0:
        return None
    return share_lokal / share_referensi

def label_ocean(lq, n, pct_populasi=None, min_n=MIN_N, min_pct=MIN_PCT):
    """
    Klasifikasi Red/Neutral/Blue Ocean murni dari Location Quotient (LQ),
    tanpa komponen pertumbuhan apa pun. Ambang 0.8/1.25 mengikuti
    konvensi standar literatur ekonomi regional (Location Quotient),
    bukan angka buatan sendiri.
    """
    if lq is None or n < min_n or (pct_populasi is not None and pct_populasi < min_pct):
        return "Data Terbatas", "gray"
    if lq < LQ_BATAS_BLUE:
        return "Blue Ocean", "blue"
    elif lq <= LQ_BATAS_RED:
        return "Neutral Ocean", "yellow"
    else:
        return "Red Ocean", "red"

def bandingkan_sektor_lintas_wilayah(d, sektor_terpilih, referensi_dki):
    """
    d               : DataFrame lengkap (seluruh wilayah)
    sektor_terpilih : salah satu dari 6 nilai sektor_makro
    referensi_dki   : hasil hitung_referensi_lq(df_all), dihitung sekali di scope modul.
    Mengembalikan klasifikasi ocean sektor terpilih di tiap wilayah, terurut LQ ascending.
    """
    rows = []
    for wilayah, sub in d.groupby('wilayah'):
        n_wilayah = len(sub)
        sektor_sub = sub[sub['sektor_makro'] == sektor_terpilih]
        n_sektor = len(sektor_sub)
        if n_sektor == 0:
            continue
        pct_lokal = n_sektor / n_wilayah * 100
        lq = hitung_lq(sub, sektor_terpilih, referensi_dki)
        label, warna = label_ocean(lq, n_sektor, pct_populasi=pct_lokal)
        rows.append({
            'wilayah': wilayah, 'n_usaha': n_sektor,
            'lq': round(lq, 2) if lq is not None else None,
            'label': label, 'warna': warna,
        })
    result = pd.DataFrame(rows).sort_values('lq', ascending=True).reset_index(drop=True)
    return result

USIA_ORDER = ['Rintisan Baru (2024)', 'Usaha Berkembang (2022-2023)', 'Usaha Lama (2020-2021)']

def build_cluster_heatmap_data(df_region, sektor_order, usia_order=USIA_ORDER):
    """Crosstab sektor_makro x kelompok_tahun_berdiri -> klaster dominan + jumlah + persentase per sel."""
    grouped = df_region.groupby(['sektor_makro', 'kelompok_tahun_berdiri'])['cluster']
    dominan = grouped.agg(lambda s: s.mode().iloc[0] if len(s) else float('nan')).rename('cluster_dominan')
    n_total = grouped.size().rename('n')
    n_dominan = grouped.agg(lambda s: s.value_counts().iloc[0] if len(s) else 0).rename('n_dominan')
    tbl = pd.concat([dominan, n_total, n_dominan], axis=1).reset_index()
    tbl['pct_dominan'] = (tbl['n_dominan'] / tbl['n'] * 100).round(1)

    pivot_cluster = tbl.pivot(index='sektor_makro', columns='kelompok_tahun_berdiri', values='cluster_dominan').reindex(index=sektor_order, columns=usia_order)
    pivot_n = tbl.pivot(index='sektor_makro', columns='kelompok_tahun_berdiri', values='n').reindex(index=sektor_order, columns=usia_order)
    pivot_pct = tbl.pivot(index='sektor_makro', columns='kelompok_tahun_berdiri', values='pct_dominan').reindex(index=sektor_order, columns=usia_order)
    return pivot_cluster, pivot_n, pivot_pct

def discrete_colorscale(colors):
    """Colorscale Plotly diskrit (band rata, bukan gradien) untuk 1 warna per id klaster."""
    n = len(colors)
    scale = []
    for i, c in enumerate(colors):
        scale.append([i / n, c])
        scale.append([(i + 1) / n, c])
    return scale

def buat_nama_segmen(top_sektor, top_usia, usia_pct, c):
    """Penamaan segmen bisnis komunikatif berbasis profil dominan aktual (bukan tebakan)."""
    if usia_pct >= 60:
        deskriptor_usia = top_usia.split(' (')[0]
        return f"{top_sektor} — Segmen {deskriptor_usia} (Klaster {c})"
    return f"{top_sektor} (Klaster {c})"

def generate_cluster_profiles(df_region, profile_region):
    """Menyusun profil klaster 2-dimensi (sektor, usia usaha) berbasis tabel profil
    final (03_kmodes_evaluation_summary.csv) yang sudah berisi kategori dominan
    deskriptif, Location Quotient (LQ), dan label Red/Neutral/Blue Ocean per klaster
    -- dihitung di notebooks/03_interpretasi_dan_penerapan_ocean.ipynb, bukan
    dihitung ulang di sini -- lalu menyematkan nama segmen ke data mentah wilayah tsb."""
    df = df_region.copy()
    cluster_profiles = {}
    label_map = {}

    profile_region = profile_region.sort_values('cluster')

    for row in profile_region.itertuples():
        c = int(row.cluster)
        nama_segmen = buat_nama_segmen(
            row.sektor_makro_dominan, row.kelompok_tahun_berdiri_dominan,
            row.kelompok_tahun_berdiri_pct, c
        )

        label_map[c] = nama_segmen
        cluster_profiles[nama_segmen] = {
            'id': c,
            'total': int(row.n),
            'pct': row.pct_n,
            'top_sektor': row.sektor_makro_dominan,
            'top_sektor_pct': row.sektor_makro_pct,
            'top_usia': row.kelompok_tahun_berdiri_dominan,
            'top_usia_pct': row.kelompok_tahun_berdiri_pct,
            # Warna mengikuti id klaster supaya sama dengan warna sel heatmap
            'border': CLUSTER_COLORS[c % len(CLUSTER_COLORS)],
            'bg': CLUSTER_BG[c % len(CLUSTER_BG)],
            'lq': row.lq if pd.notna(row.lq) else None,
            'label_ocean': row.label_ocean,
            'warna_ocean': row.warna_ocean,
        }

    df['nama_segmen'] = df['cluster'].map(label_map)
    return df, cluster_profiles

def tampilkan_peta_sektor(hasil_sektor):
    """Peta choropleth 6 wilayah DKI (Folium/Leaflet): warna poligon = label Ocean sektor
    terpilih (LQ per wilayah administrasi). Klaster tidak dipetakan karena tidak bersifat
    spasial. Jika GeoJSON/folium tidak tersedia, cukup tampilkan st.info (tidak error)."""
    if not FOLIUM_TERSEDIA:
        st.info("Peta tidak ditampilkan karena paket `folium` dan `streamlit-folium` belum terpasang. Jalankan `pip install -r requirements.txt`.")
        return

    geojson, pesan_error = load_geojson()
    if pesan_error is not None:
        st.info(
            f"🗺️ Peta tidak ditampilkan: {pesan_error} Taruh berkas GeoJSON batas 6 wilayah DKI Jakarta "
            f"(dengan properti `CC_2` berisi kode wilayah BPS: {', '.join(KODE_WILAYAH)}) di `data/geo/{GEOJSON_FILENAME}` "
            f"-- lihat `data/geo/README.md`. Tabel dan grafik batang di bawah tetap berlaku."
        )
        return

    # Sisipkan atribut sektor terpilih ke properti tiap poligon (untuk warna & tooltip)
    data_wilayah = hasil_sektor.set_index('wilayah').to_dict('index')
    geojson = copy.deepcopy(geojson)
    for ft in geojson['features']:
        prop = ft['properties']
        baris = data_wilayah.get(prop['wilayah'])
        lq = baris['lq'] if baris is not None else None
        prop['n_usaha'] = f"{int(baris['n_usaha']):,}" if baris is not None else "0"
        prop['lq_teks'] = f"{lq:.2f}" if lq is not None and pd.notna(lq) else "N/A"
        prop['label'] = baris['label'] if baris is not None else "Data Terbatas"
        prop['warna'] = baris['warna'] if baris is not None else "gray"
        # Nama pendek untuk tooltip, mis. "KOTA ADM. JAKARTA SELATAN" -> "Jakarta Selatan"
        prop['nama_pendek'] = prop['wilayah'].replace("KOTA ADM. ", "").replace("KAB. ADM. ", "").title()

    # Batas tampilan awal = batas gabungan seluruh poligon (termasuk Kep. Seribu di utara)
    semua_titik = [
        titik
        for ft in geojson['features']
        for poligon in (ft['geometry']['coordinates'] if ft['geometry']['type'] == 'MultiPolygon' else [ft['geometry']['coordinates']])
        for cincin in poligon
        for titik in cincin
    ]
    batas_peta = [
        [min(t[1] for t in semua_titik), min(t[0] for t in semua_titik)],
        [max(t[1] for t in semua_titik), max(t[0] for t in semua_titik)],
    ]

    # Tanpa basemap (tiles=None): hanya poligon dari GeoJSON, latar transparen mengikuti tema halaman
    peta = folium.Map(
        tiles=None, attribution_control=False, zoom_control=False,
        min_zoom=8, max_zoom=14, zoom_snap=0.1,
    )
    peta.get_root().header.add_child(folium.Element(
        "<style>html, body, .folium-map, .leaflet-container { background: transparent !important; }</style>"
    ))
    folium.GeoJson(
        geojson,
        name="Wilayah DKI Jakarta",
        # Pulau-pulau Kep. Seribu sangat kecil pada skala ini, jadi diberi garis tepi tebal
        # berwarna sama dengan isinya supaya tetap terlihat sebagai titik-titik pulau
        style_function=lambda ft: {
            'fillColor': BADGE_WARNA[ft['properties']['warna']],
            'color': BADGE_WARNA[ft['properties']['warna']] if ft['properties']['wilayah'] == WILAYAH_KEP_SERIBU else '#FFFFFF',
            'weight': 4 if ft['properties']['wilayah'] == WILAYAH_KEP_SERIBU else 1.2,
            'fillOpacity': 0.85,
        },
        highlight_function=lambda ft: {'fillOpacity': 1.0, 'weight': 6 if ft['properties']['wilayah'] == WILAYAH_KEP_SERIBU else 3},
        # Tooltip ringkas (nama pendek, label singkat, font kecil) supaya tidak terpotong tepi peta
        tooltip=folium.GeoJsonTooltip(
            fields=['nama_pendek', 'n_usaha', 'lq_teks', 'label'],
            aliases=['Wilayah', 'Usaha', 'LQ', 'Label'],
            sticky=True,
            style=(
                "background-color: #FFFFFF; color: #111827; font-size: 11.5px; "
                "padding: 4px 8px; border: 1px solid #D1D5DB; border-radius: 6px; "
                "box-shadow: 0 2px 6px rgba(0,0,0,0.25); white-space: nowrap;"
            ),
        ),
    ).add_to(peta)

    # Label "Kep. Seribu" di samping kanan gugusan pulau (titik-titik kecil sulit dikenali tanpa label);
    # posisinya = rata-rata titik pulau, jadi mengikuti pusat gugusan, bukan pulau terluar
    ft_seribu = next((ft for ft in geojson['features'] if ft['properties']['wilayah'] == WILAYAH_KEP_SERIBU), None)
    if ft_seribu is not None:
        titik_seribu = [t for poligon in ft_seribu['geometry']['coordinates'] for cincin in poligon for t in cincin]
        lat_tengah = sum(t[1] for t in titik_seribu) / len(titik_seribu)
        lon_tengah = sum(t[0] for t in titik_seribu) / len(titik_seribu)
        folium.Marker(
            [lat_tengah, lon_tengah],
            icon=folium.DivIcon(
                icon_size=(80, 20), icon_anchor=(-24, 10),
                html=(
                    '<div style="font-size:11px;font-weight:600;color:#111827;background:rgba(255,255,255,0.9);'
                    'padding:1px 6px;border-radius:8px;white-space:nowrap;display:inline-block;">Kep. Seribu</div>'
                ),
            ),
        ).add_to(peta)

    peta.fit_bounds(batas_peta, padding=(14, 14))

    # Peta tinggi-ramping (bentuk DKI + Kep. Seribu memanjang ke utara), lebar tetap di kolom tengah
    _, kolom_peta, _ = st.columns([1, 1.1, 1])
    with kolom_peta:
        st_folium(peta, width=360, height=560, returned_objects=[], key="peta_sektor")

        legend_html = " &nbsp; ".join(
            f'<span style="white-space:nowrap;"><span style="display:inline-block;width:11px;height:11px;background:{BADGE_WARNA[w]};border-radius:2px;vertical-align:middle;"></span> {nama}</span>'
            for w, nama in [('blue', 'Blue Ocean (LQ < 0.8)'), ('yellow', 'Neutral Ocean (0.8–1.25)'),
                            ('red', 'Red Ocean (LQ > 1.25)'), ('gray', 'Data Terbatas')]
        )
        st.markdown(f'<div style="font-size:12px;text-align:center;line-height:1.8;">{legend_html}</div>', unsafe_allow_html=True)
    st.caption("Peta menampilkan konsentrasi sektor (LQ) per wilayah administrasi, bukan sebaran klaster. Kep. Seribu (pulau-pulau kecil di bagian atas peta) dapat diperbesar dengan scroll.")

@st.cache_data(show_spinner="Menyiapkan berkas Excel...")
def buat_excel_segmentasi(df_data, df_ringkasan):
    """Menyusun berkas Excel (.xlsx) 2 sheet -- 'Data Segmentasi' (daftar usaha) dan
    'Ringkasan Segmen' (profil tiap klaster) -- dengan header tebal yang dibekukan,
    filter di tiap kolom, dan lebar kolom menyesuaikan isi. Di-cache per wilayah."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for nama_sheet, d in [("Data Segmentasi", df_data), ("Ringkasan Segmen", df_ringkasan)]:
            d.to_excel(writer, sheet_name=nama_sheet, index=False)
            ws = writer.sheets[nama_sheet]

            for sel in ws[1]:
                sel.font = Font(bold=True, color="FFFFFF")
                sel.fill = PatternFill("solid", fgColor="1E3A8A")
                sel.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            # Lebar kolom = teks terpanjang (header atau isi), dibatasi agar tidak terlalu lebar
            for i, kolom in enumerate(d.columns, start=1):
                panjang_isi = d[kolom].astype(str).str.len().max() if len(d) else 0
                ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(kolom)), panjang_isi) + 3, 60)
    return buffer.getvalue()

# Memuat data
df_all = load_kmodes_data()
df_elbow_all = load_elbow_data()
df_k_opt_all = load_k_optimal()
df_eval_all = load_evaluation_summary()

if df_all is None:
    st.error("⚠️ Berkas `02_umkm_kmodes_clustered.csv` tidak ditemukan! Jalankan `notebooks/02_kmodes_partitioning.ipynb` terlebih dahulu.")
    st.stop()

if df_eval_all is None:
    st.error("⚠️ Berkas `03_kmodes_evaluation_summary.csv` tidak ditemukan! Jalankan `notebooks/03_interpretasi_dan_penerapan_ocean.ipynb` terlebih dahulu.")
    st.stop()

# Referensi distribusi sektor se-DKI Jakarta untuk Location Quotient (LQ) --
# dihitung sekali di sini, dioper ke semua fungsi klasifikasi Ocean di bawah.
REFERENSI_LQ = hitung_referensi_lq(df_all)

# ==========================================
# 4. SIDEBAR: FILTER WILAYAH & PENCARIAN
# ==========================================
st.sidebar.header("📍 Filter Wilayah")

daftar_wilayah = sorted(df_all['wilayah'].unique().tolist())
selected_wilayah = st.sidebar.selectbox("Pilih Kota/Kabupaten Administrasi:", daftar_wilayah)

df_wilayah_raw = df_all[df_all['wilayah'] == selected_wilayah].copy()
profile_wilayah = df_eval_all[df_eval_all['wilayah'] == selected_wilayah].copy()
df_wilayah, profiles = generate_cluster_profiles(df_wilayah_raw, profile_wilayah)

elbow_wilayah = df_elbow_all[df_elbow_all['wilayah'] == selected_wilayah].sort_values('k') if df_elbow_all is not None else None

# k wilayah terpilih: dari 02c_k_optimal.csv; kalau berkas belum ada, pakai jumlah klaster di dataset
k_opt_wilayah = None
if df_k_opt_all is not None:
    baris_k = df_k_opt_all[df_k_opt_all['wilayah'] == selected_wilayah]
    if not baris_k.empty:
        k_opt_wilayah = baris_k.iloc[0]
k_wilayah = int(k_opt_wilayah['k_optimal']) if k_opt_wilayah is not None else int(df_wilayah_raw['cluster'].nunique())

st.sidebar.markdown("---")
st.sidebar.header("🔍 Cari Komoditas Usaha")
search_query = st.sidebar.text_input("Ketik jenis usaha (misal: Sablon, Ikan, Fesyen):", "")

if search_query:
    df_search = df_wilayah[
        df_wilayah['jenis_usaha'].str.contains(search_query, case=False, na=False) |
        df_wilayah['sektor_makro'].str.contains(search_query, case=False, na=False)
    ]
    st.sidebar.info(f"Ditemukan {len(df_search):,} data dengan kata kunci '{search_query}'.")

st.sidebar.markdown("---")
st.sidebar.caption("💡 *Sistem Segmentasi UMKM DKI Jakarta - Algoritma K-Modes Clustering (2 Fitur: Sektor, Usia Usaha).*")

# ==========================================
# 5. JUDUL APLIKASI & TAB UTAMA
# ==========================================
st.markdown('<div class="main-title">Peta Segmentasi & Peluang Usaha UMKM DKI Jakarta</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-title">Segmentasi Data Terbuka Berbasis K-Modes Clustering (Periode 2024) & Penerapan Red/Blue Ocean Menggunakan Location Quotient</div>', unsafe_allow_html=True)

tab_wilayah, tab_sektor, tab_k = st.tabs([
    "📍 Analisis per Wilayah", "📊 Perbandingan Sektor Lintas Wilayah", "📉 Evaluasi Jumlah Klaster (k)"
])

# ==========================================
# TAB 1: ANALISIS PER WILAYAH (alur yang sudah ada)
# ==========================================
with tab_wilayah:

    st.markdown(f"### {selected_wilayah}")

    kpi1, kpi2, kpi3, kpi4 = st.columns(4)
    with kpi1:
        st.metric("Total UMKM Terdata", f"{len(df_wilayah):,} Unit")
    with kpi2:
        st.metric("Ragam Jenis Usaha", f"{df_wilayah['jenis_usaha'].nunique():,} Kategori")
    with kpi3:
        st.metric("Sektor Terbesar", df_wilayah['sektor_makro'].mode()[0])
    with kpi4:
        st.metric("Jumlah Segmen", f"{k_wilayah} Klaster", help="Jumlah segmen (k) UMKM di wilayah ini, dihitung dengan aturan titik terjauh dari garis lurus pada kurva Elbow k=2..6 (lihat tab Evaluasi Jumlah Klaster dan notebooks/02_kmodes_partitioning.ipynb).")

    st.markdown("---")

    # ==========================================
    # KARTU PROFIL KLASTER MULTIDIMENSI
    # ==========================================
    st.subheader("💡 Profil Klaster (Sektor × Usia Usaha)")
    profiles_terurut = sorted(profiles.items(), key=lambda item: item[1]['pct'], reverse=True)
    cols_kartu = st.columns(len(profiles_terurut))

    for idx, (c_name, info) in enumerate(profiles_terurut):
        badge_color = BADGE_WARNA[info['warna_ocean']]
        badge_emoji = BADGE_EMOJI[info['warna_ocean']]
        lq_text = f"{info['lq']:.2f}" if info['lq'] is not None else "N/A"
        with cols_kartu[idx]:
            st.markdown(f"""
                <div style="
                    background-color: {info['bg']};
                    border-left: 5px solid {info['border']};
                    padding: 14px;
                    border-radius: 8px;
                    min-height: 260px;
                ">
                    <h4 style="margin: 0 0 6px 0; color: #111827; font-size: 14.5px;">{c_name}</h4>
                    <p style="font-size: 12.5px; margin: 0; color: #374151; line-height: 1.6;">
                        <b>Populasi:</b> {info['total']:,} UMKM ({info['pct']:.1f}%)<br>
                        <b>Sektor Dominan:</b> {info['top_sektor']} ({info['top_sektor_pct']:.1f}%)<br>
                        <b>Usia Usaha Dominan:</b> {info['top_usia']} ({info['top_usia_pct']:.1f}%)<br>
                    </p>
                    <div style="margin-top: 10px;">
                        <span title="{OCEAN_TOOLTIP}" style="background-color: {badge_color}; color: white; padding: 3px 10px; border-radius: 12px; font-size: 11px; font-weight: 600; cursor: help;">
                            LQ: {lq_text} {badge_emoji} {info['label_ocean']}
                        </span>
                    </div>
                </div>
            """, unsafe_allow_html=True)

    # LQ klaster = LQ sektor dominannya di wilayah ini, jadi sektor dominan sama -> badge Ocean identik
    if len({info['top_sektor'] for info in profiles.values()}) == 1:
        st.caption("Sektor dominan sama di seluruh klaster, perbandingan Ocean antar-sektor ada di tab Perbandingan Sektor")

    # ==========================================
    # PETA KLASTER: HEATMAP SEKTOR x USIA USAHA
    # ==========================================
    st.subheader("🗺️ Peta Klaster: Sektor × Usia Usaha")
    st.caption("Setiap sel menunjukkan klaster yang mendominasi kombinasi sektor × usia usaha tsb, beserta jumlah dan persentase UMKM dari kombinasi itu yang benar-benar berada di klaster dominan tersebut.")

    sektor_order_wilayah = sorted(df_wilayah['sektor_makro'].unique().tolist())
    pivot_cluster, pivot_n, pivot_pct = build_cluster_heatmap_data(df_wilayah, sektor_order_wilayah)

    text_matrix = []
    for r in range(pivot_cluster.shape[0]):
        row_text = []
        for c in range(pivot_cluster.shape[1]):
            if pd.isna(pivot_cluster.iat[r, c]):
                row_text.append("Tidak ada data")
            else:
                row_text.append(
                    f"Klaster {int(pivot_cluster.iat[r, c])} • {int(pivot_n.iat[r, c]):,} UMKM ({pivot_pct.iat[r, c]:.0f}%)"
                )
        text_matrix.append(row_text)

    fig_heatmap = go.Figure(go.Heatmap(
        z=pivot_cluster.values, x=[u.split(' (')[0] for u in USIA_ORDER], y=sektor_order_wilayah,
        colorscale=discrete_colorscale(CLUSTER_COLORS[:k_wilayah]), zmin=-0.5, zmax=k_wilayah - 0.5,
        text=text_matrix, texttemplate="%{text}", textfont={"size": 11},
        hovertemplate="Sektor: %{y}<br>Usia Usaha: %{x}<br>%{text}<extra></extra>",
        xgap=4, ygap=4,
        colorbar=dict(title="Klaster", tickmode="array", tickvals=list(range(k_wilayah)), ticktext=[f"Klaster {i}" for i in range(k_wilayah)]),
    ))
    fig_heatmap.update_layout(margin=dict(t=20, b=10, l=10, r=10), height=380)
    st.plotly_chart(fig_heatmap, width='stretch')

    # ==========================================
    # VISUALISASI KOMPOSISI & KOMODITAS
    # ==========================================
    col_viz1, col_viz2 = st.columns([1, 1.3])

    with col_viz1:
        st.subheader("Komposisi Segmen UMKM")
        cluster_counts = df_wilayah['nama_segmen'].value_counts().reset_index()
        cluster_counts.columns = ['Segmen', 'Jumlah UMKM']
        # Urut menurut id klaster supaya urutan legend sama dengan kartu profil & heatmap
        cluster_counts['id'] = cluster_counts['Segmen'].map({nama: info['id'] for nama, info in profiles.items()})
        cluster_counts = cluster_counts.sort_values('id')

        fig_pie = px.pie(
            cluster_counts,
            values='Jumlah UMKM',
            names='Segmen',
            hole=0.5,
            color='Segmen',
            color_discrete_map={nama: info['border'] for nama, info in profiles.items()},
            category_orders={'Segmen': cluster_counts['Segmen'].tolist()},
        )
        # Persentase di dalam irisan, legend di bawah donut (nama segmen panjang tidak memakan lebar donut);
        # tinggi 430 px menyamai kolom kanan (selectbox + grafik batang 340 px)
        fig_pie.update_traces(
            sort=False, texttemplate='%{percent:.1%}', textposition='inside', insidetextorientation='horizontal',
            textfont=dict(size=13, color='white'),
            hovertemplate="%{label}<br>%{value:,} UMKM (%{percent})<extra></extra>",
        )
        fig_pie.update_layout(
            height=430, margin=dict(t=10, b=10, l=10, r=10),
            legend=dict(orientation='h', yanchor='top', y=-0.02, xanchor='center', x=0.5, font=dict(size=11)),
            uniformtext=dict(minsize=10, mode='hide'),
        )
        st.plotly_chart(fig_pie, width='stretch')

    with col_viz2:
        st.subheader("Komoditas Terbanyak per Segmen")
        pilihan_segmen_chart = st.selectbox("Pilih Segmen untuk Ditinjau:", list(profiles.keys()), key=f"segmen_select_{selected_wilayah}")

        df_sub = df_wilayah[df_wilayah['nama_segmen'] == pilihan_segmen_chart]
        top_jenis = df_sub['jenis_usaha'].value_counts().head(8).reset_index()
        top_jenis.columns = ['Jenis Usaha', 'Jumlah Unit']

        fig_bar = px.bar(
            top_jenis,
            x='Jumlah Unit',
            y='Jenis Usaha',
            orientation='h',
            text='Jumlah Unit',
            color='Jumlah Unit',
            color_continuous_scale='Blues'
        )
        fig_bar.update_layout(yaxis=dict(autorange="reversed"), margin=dict(t=20, b=10, l=10, r=10), height=340, coloraxis_showscale=False)
        st.plotly_chart(fig_bar, width='stretch')

    # ==========================================
    # TABEL EKSPLORASI DAFTAR USAHA
    # ==========================================
    st.markdown("---")
    st.subheader("📋 Eksplorasi Daftar Usaha Tersegmentasi")

    kolom_tampil = ['jenis_usaha', 'sektor_makro', 'kelompok_tahun_berdiri', 'nama_segmen']

    if search_query:
        st.markdown(f"**Hasil Pencarian untuk:** `{search_query}`")
        st.dataframe(df_search[kolom_tampil + ['wilayah']], width='stretch')
    else:
        list_names = list(profiles.keys())
        tabs = st.tabs([f"📂 {name}" for name in list_names])

        for idx_tab, tab_obj in enumerate(tabs):
            target_name = list_names[idx_tab]
            with tab_obj:
                df_tab_data = df_wilayah[df_wilayah['nama_segmen'] == target_name]

                rekap_usaha = df_tab_data.groupby(['jenis_usaha', 'sektor_makro']).agg(
                    Jumlah_Usaha=('jenis_usaha', 'count'),
                    Usia_Usaha_Utama=('kelompok_tahun_berdiri', lambda x: x.mode()[0] if not x.empty else '-')
                ).reset_index().sort_values(by='Jumlah_Usaha', ascending=False)

                rekap_usaha['Porsi di Segmen (%)'] = (
                    (rekap_usaha['Jumlah_Usaha'] / len(df_tab_data)) * 100
                ).round(2)

                st.dataframe(
                    rekap_usaha,
                    column_config={
                        "jenis_usaha": st.column_config.TextColumn("Jenis Usaha", width="large"),
                        "sektor_makro": st.column_config.TextColumn("Sektor Makro", width="medium"),
                        "Jumlah_Usaha": st.column_config.NumberColumn("Jumlah Unit", format="%d unit"),
                        "Porsi di Segmen (%)": st.column_config.NumberColumn("Porsi di Segmen", format="%.2f %%"),
                        "Usia_Usaha_Utama": "Usia Usaha"
                    },
                    width='stretch',
                    hide_index=True
                )

    # ==========================================
    # UNDUH DATA
    # ==========================================
    data_export = df_wilayah[['wilayah'] + kolom_tampil].rename(columns={
        'wilayah': 'Wilayah', 'jenis_usaha': 'Jenis Usaha', 'sektor_makro': 'Sektor Makro',
        'kelompok_tahun_berdiri': 'Usia Usaha', 'nama_segmen': 'Segmen',
    })
    ringkasan_export = pd.DataFrame([
        {
            'Klaster': info['id'], 'Segmen': nama, 'Jumlah UMKM': info['total'],
            'Porsi Populasi (%)': round(info['pct'], 2),
            'Sektor Dominan': info['top_sektor'], 'Porsi Sektor Dominan (%)': round(info['top_sektor_pct'], 2),
            'Usia Usaha Dominan': info['top_usia'], 'Porsi Usia Dominan (%)': round(info['top_usia_pct'], 2),
            'LQ': round(info['lq'], 3) if info['lq'] is not None else None, 'Label Ocean': info['label_ocean'],
        }
        for nama, info in sorted(profiles.items(), key=lambda item: item[1]['id'])
    ])
    xlsx_export = buat_excel_segmentasi(data_export, ringkasan_export)
    nama_file = selected_wilayah.lower().replace('.', '').replace(' ', '_')
    st.download_button(
        label=f"📥 Unduh Data Segmentasi {selected_wilayah} (Excel)",
        data=xlsx_export,
        file_name=f"segmentasi_{nama_file}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# ==========================================
# TAB 2: PERBANDINGAN SEKTOR LINTAS WILAYAH
# ==========================================
with tab_sektor:
    st.subheader("📊 Perbandingan 1 Sektor Lintas 6 Wilayah")
    st.caption("Mode analisis lain: pilih 1 sektor usaha, lihat klasifikasi Red/Neutral/Blue Ocean-nya di seluruh 6 wilayah DKI Jakarta -- independen dari filter wilayah di sidebar.")

    daftar_sektor = sorted(df_all['sektor_makro'].unique().tolist())
    sektor_terpilih = st.selectbox("Pilih Sektor Makro:", daftar_sektor, key="sektor_lintas_wilayah_select")

    hasil_sektor = bandingkan_sektor_lintas_wilayah(df_all, sektor_terpilih, REFERENSI_LQ)

    if hasil_sektor.empty:
        st.warning(f"Tidak ada data untuk sektor '{sektor_terpilih}' di wilayah manapun.")
    else:
        lq_valid = hasil_sektor[hasil_sektor['lq'].notna()]
        if lq_valid.empty:
            st.info(f"Location Quotient (LQ) sektor **{sektor_terpilih}** tidak dapat dihitung untuk wilayah mana pun.")
        else:
            tertinggi = lq_valid.loc[lq_valid['lq'].idxmax()]
            terendah = lq_valid.loc[lq_valid['lq'].idxmin()]
            st.info(
                f"Sektor **{sektor_terpilih}** paling terkonsentrasi di **{tertinggi['wilayah']}** "
                f"(LQ {tertinggi['lq']:.2f}, {tertinggi['label']}) dan paling longgar di **{terendah['wilayah']}** "
                f"(LQ {terendah['lq']:.2f}, {terendah['label']})."
            )

        tampilkan_peta_sektor(hasil_sektor)

        tabel_sektor = hasil_sektor[[
            'wilayah', 'n_usaha', 'lq', 'label'
        ]].copy()
        tabel_sektor.columns = [
            'Wilayah', 'Jumlah Usaha', 'Location Quotient (LQ)', 'Label'
        ]
        st.dataframe(
            tabel_sektor,
            column_config={
                "Wilayah": st.column_config.TextColumn("Wilayah", width="medium"),
                "Location Quotient (LQ)": st.column_config.NumberColumn("Location Quotient (LQ)", format="%.2f", help=LQ_TOOLTIP),
            },
            width='stretch',
            hide_index=True,
        )

        fig_sektor = px.bar(
            hasil_sektor,
            x='lq', y='wilayah', orientation='h',
            text='lq', color='warna',
            color_discrete_map=BADGE_WARNA,
        )
        fig_sektor.update_traces(texttemplate='%{text:.2f}', textposition='outside')
        fig_sektor.add_vline(x=1.0, line_dash="dash", line_color="#6B7280", annotation_text="Rata-rata DKI (LQ=1.0)")
        fig_sektor.update_layout(
            xaxis_title="Location Quotient (LQ)", yaxis_title="",
            showlegend=False, margin=dict(t=20, b=10, l=10, r=30), height=340,
            xaxis_range=[0, max(1.5, hasil_sektor['lq'].max(skipna=True) * 1.15 if hasil_sektor['lq'].notna().any() else 1.5)]
        )
        st.plotly_chart(fig_sektor, width='stretch')

        with st.expander("Apa itu Red/Blue Ocean?", expanded=False):
            st.caption(OCEAN_TOOLTIP)

# ==========================================
# TAB 3: EVALUASI JUMLAH KLASTER (k)
# ==========================================
with tab_k:
    st.subheader("📉 Evaluasi Jumlah Klaster (k) per Wilayah")

    if df_k_opt_all is None:
        st.info("Berkas `02c_k_optimal.csv` belum tersedia. Jalankan `notebooks/02_kmodes_partitioning.ipynb` terlebih dahulu.")
    else:
        tabel_k = df_k_opt_all[['wilayah', 'k_optimal', 'kekuatan_siku', 'catatan']].copy()
        tabel_k.columns = ['Wilayah', 'k Optimal', 'Kekuatan Siku', 'Catatan']
        st.dataframe(
            tabel_k,
            column_config={
                "Wilayah": st.column_config.TextColumn("Wilayah", width="medium"),
                "Kekuatan Siku": st.column_config.NumberColumn(
                    "Kekuatan Siku (d)", format="%.4f",
                    help="Jarak titik k optimal ke garis lurus pada kurva Elbow ternormalisasi. Makin besar, makin tegas sikunya."
                ),
                "Catatan": st.column_config.TextColumn("Catatan", width="large"),
            },
            width='stretch',
            hide_index=True,
        )

    st.markdown(f"**Kurva Elbow — {selected_wilayah}**")
    if elbow_wilayah is not None and not elbow_wilayah.empty:
        fig_elbow = go.Figure()
        fig_elbow.add_trace(go.Scatter(
            x=elbow_wilayah['k'], y=elbow_wilayah['cost'], name='Cost',
            mode='lines+markers', line=dict(color='#2563EB', width=3), marker=dict(size=9)
        ))
        # Garis lurus penghubung titik pertama-terakhir (acuan aturan titik terjauh)
        fig_elbow.add_trace(go.Scatter(
            x=elbow_wilayah['k'].iloc[[0, -1]], y=elbow_wilayah['cost'].iloc[[0, -1]], name='Garis acuan',
            mode='lines', line=dict(color='#9CA3AF', width=1.5, dash='dot')
        ))
        teks_k = f"k optimal = {k_wilayah}"
        if k_opt_wilayah is not None:
            teks_k += f" (d = {k_opt_wilayah['kekuatan_siku']:.3f})"
        fig_elbow.add_vline(x=k_wilayah, line_dash="dash", line_color="#EF4444", annotation_text=teks_k)
        fig_elbow.update_layout(
            xaxis_title="Jumlah Klaster (k)", yaxis_title="Nilai Cost", showlegend=False,
            xaxis=dict(tickmode='array', tickvals=elbow_wilayah['k'].tolist()),
            margin=dict(t=30, b=10, l=10, r=10), height=340
        )
        st.plotly_chart(fig_elbow, width='stretch')
    else:
        st.info("Data kurva biaya untuk wilayah ini belum tersedia.")

    if df_k_opt_all is not None:
        siku_lemah = df_k_opt_all[df_k_opt_all['wilayah'].isin(['KOTA ADM. JAKARTA BARAT', 'KOTA ADM. JAKARTA UTARA'])]
        rincian = "; ".join(
            f"{r.wilayah}: d = {r.kekuatan_siku:.3f} pada k={r.k_optimal}, hanya berselisih {r.selisih_d:.4f} dari k={r.k_pembanding}"
            for r in siku_lemah.itertuples()
        )

st.markdown("---")
st.caption("Aplikasi Analitik Segmentasi UMKM DKI Jakarta | K-Modes Clustering (Sektor Makro × Usia Usaha) | Penentuan k: Kurva Elbow (Titik Terjauh dari Garis Lurus)")
